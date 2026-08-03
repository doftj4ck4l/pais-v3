#!/usr/bin/env python3
"""Create and read .xlsx workbooks.

Runs inside the sandbox container with the user's workspace as the working
directory, so a relative output path lands where the console can download it.
"""
import json
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_WIDTH = 60


def fail(message):
    print(f"error: {message}", file=sys.stderr)
    sys.exit(1)


def create(spec):
    path = spec.get("path")
    if not path:
        fail('"path" is required, e.g. "report.xlsx"')

    sheets = spec.get("sheets") or []
    if not sheets:
        fail('"sheets" must contain at least one sheet')

    book = Workbook()
    book.remove(book.active)
    summary = []

    for index, sheet in enumerate(sheets):
        rows = sheet.get("rows") or []
        title = str(sheet.get("name") or f"Sheet{index + 1}")[:31]
        ws = book.create_sheet(title=title)

        for row in rows:
            ws.append(list(row) if isinstance(row, (list, tuple)) else [row])

        if sheet.get("header") and rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"

        # Size columns to their content so the file is readable when opened,
        # rather than a wall of ### in numeric columns.
        for column in range(1, ws.max_column + 1):
            longest = 0
            for cell in ws[get_column_letter(column)]:
                if cell.value is not None:
                    longest = max(longest, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(column)].width = min(longest + 2, MAX_WIDTH)

        summary.append(f"  {title}: {ws.max_row} row(s) x {ws.max_column} column(s)")

    book.save(path)
    print(f"wrote {path}")
    print("\n".join(summary))


def read(spec):
    path = spec.get("path")
    if not path:
        fail('"path" is required')
    max_rows = int(spec.get("maxRows") or 50)

    try:
        book = load_workbook(path, data_only=True)
    except FileNotFoundError:
        fail(f"no such file: {path}")

    for ws in book.worksheets:
        print(f"# {ws.title} ({ws.max_row} rows x {ws.max_column} cols)")
        for row in ws.iter_rows(max_row=min(ws.max_row, max_rows), values_only=True):
            print(" | ".join("" if v is None else str(v) for v in row))
        if ws.max_row > max_rows:
            print(f"... {ws.max_row - max_rows} more row(s)")
        print()


def main():
    if len(sys.argv) < 2:
        fail('expected a JSON argument, e.g. {"op":"create","path":"a.xlsx","sheets":[...]}')
    try:
        spec = json.loads(sys.argv[1])
    except json.JSONDecodeError as e:
        fail(f"argument is not valid JSON: {e}")

    op = spec.get("op")
    if op == "create":
        create(spec)
    elif op == "read":
        read(spec)
    else:
        fail('"op" must be "create" or "read"')


if __name__ == "__main__":
    main()
